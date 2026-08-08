"""Generate CRM-ready internal-code mappings and reviewed catalogue updates."""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_CATALOGUE = "San_pham.xlsx"
DEFAULT_INVENTORY = "Tong_hop_ton_kho.xlsx"
MAPPING_HEADERS = ("Mã báo giá", "Mã nội bộ")
CATALOGUE_FIELDS = ("name", "size", "packing", "packing_m2", "packing_pcs")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
WRAP = Alignment(vertical="top", wrap_text=True)


@dataclass(frozen=True)
class InventoryRow:
    source_row: int
    internal_code: str
    item_name: str
    description: str
    stock_location: str
    quantity: float | None


@dataclass
class MatchResult:
    row: InventoryRow
    product_code: str = ""
    rule: str = "unmatched"
    reason: str = ""
    name: str = ""
    size: str = ""
    packing: str = ""
    packing_m2: float | None = None
    packing_pcs: float | None = None


@dataclass
class MappingCandidate:
    internal_code: str
    product_code: str = ""
    status: str = "unmatched"
    rows: list[MatchResult] = field(default_factory=list)
    reason: str = ""


def text(value: object) -> str:
    return str(value if value is not None else "").strip()


def normalized_header(value: object) -> str:
    value = unicodedata.normalize("NFD", text(value).lower())
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", value).strip()


def normalized_code(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", text(value).upper())


def code_key(value: object) -> str:
    return text(value).casefold()


def candidates(raw_code: str) -> list[str]:
    """Return code variants from least to most destructive."""
    code = text(raw_code).upper()
    seen: set[str] = set()
    result: list[str] = []

    def add(value: str) -> None:
        value = text(value)
        if value and value not in seen:
            seen.add(value)
            result.append(value)

    add(code)
    add(normalized_code(code))
    prefix_patterns = (r"^$", r"^\d+-", r"^\d(?=[A-Z0-9])", r"^[A-Z](?=\d)")
    suffix_patterns = (r"$", r"-HN$", r"-\d+$", r"-[A-Z]$", r"[A-Z]$")
    combinations = sorted(
        ((prefix_index + suffix_index, prefix, suffix)
         for prefix_index, prefix in enumerate(prefix_patterns)
         for suffix_index, suffix in enumerate(suffix_patterns)),
        key=lambda item: item[0],
    )
    for _, prefix, suffix in combinations:
        stripped = re.sub(prefix, "", code, count=1)
        stripped = re.sub(suffix, "", stripped, count=1)
        add(stripped)
        add(normalized_code(stripped))
    return result


def extract_alias_codes(description: str) -> list[str]:
    if not text(description):
        return []
    aliases: list[str] = []

    def add_tokens(value: str) -> None:
        for token in re.split(r"\s*/\s*", value):
            token = token.strip().strip("/").strip()
            if token:
                aliases.append(token)

    prefix = re.match(r"^([^()]+)", description)
    if prefix:
        add_tokens(prefix.group(1))
    for group in re.finditer(r"\(([^()]+)\)", description):
        add_tokens(group.group(1))
    return aliases


def extract_name_size_packing(item_name: str, internal_code: str) -> tuple[str, str, str]:
    content = text(item_name)
    if not content:
        return "", "", ""
    packing = ""
    last_parentheses = re.search(r"\(([^()]*)\)\s*$", content)
    without_packing = content
    if last_parentheses:
        packing = text(last_parentheses.group(1))
        without_packing = content[:last_parentheses.start()].rstrip()
    position = without_packing.upper().find(text(internal_code).upper())
    if position >= 0:
        return (
            without_packing[:position].rstrip(" ,").strip(),
            without_packing[position + len(internal_code):].lstrip(" ,").strip(),
            packing,
        )
    if "," in without_packing:
        name, size = without_packing.split(",", 1)
        return name.strip(), size.strip(), packing
    return without_packing, "", packing


def parse_number(value: object) -> float | None:
    if value is None or text(value) in {"", "-", "—"}:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    raw = text(value).replace(" ", "").replace("%", "")
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", raw):
        raw = raw.replace(".", "").replace(",", "")
    else:
        raw = raw.replace(",", ".") if raw.count(",") == 1 and "." not in raw else raw.replace(",", "")
    try:
        return float(raw)
    except ValueError:
        return None


def parse_packaging_number(value: str) -> float | None:
    """Read packaging decimals such as 1.202m² as 1.202, never as 1202."""
    raw = text(value).replace(" ", "")
    if not re.fullmatch(r"\d+(?:[.,]\d+)?", raw):
        return None
    try:
        return float(raw.replace(",", "."))
    except ValueError:
        return None


def parse_packaging(packing: str) -> tuple[float | None, float | None]:
    """Extract pieces/pack and m2/pack without inventing a unit."""
    content = text(packing)
    if not content:
        return None, None
    pieces: float | None = None
    area: float | None = None
    for part in re.split(r"/", content):
        part = text(part)
        match = re.match(r"^([\d.,]+)\s*(.*)$", part)
        if not match:
            continue
        number = parse_packaging_number(match.group(1))
        unit = normalized_header(match.group(2))
        if number is None:
            continue
        if "m2" in unit or "m²" in text(match.group(2)).lower() or "㎡" in match.group(2):
            area = number
        elif pieces is None:
            pieces = number
    return pieces, area


def header_positions(row: Iterable[object]) -> dict[str, int]:
    return {normalized_header(value): index for index, value in enumerate(row) if text(value)}


def locate_headers(ws, required: set[str], max_rows: int = 20) -> tuple[int, dict[str, int]]:
    for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_number > max_rows:
            break
        positions = header_positions(row)
        if required.issubset(positions):
            return row_number, positions
    expected = ", ".join(sorted(required))
    raise ValueError(f"Không tìm thấy header bắt buộc: {expected}")


def read_catalogue(path: Path) -> dict[str, dict[str, object]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for worksheet in workbook.worksheets:
            try:
                header_row, positions = locate_headers(worksheet, {"code"})
            except ValueError:
                continue
            products: dict[str, dict[str, object]] = {}
            for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                code = text(values[positions["code"]] if positions["code"] < len(values) else "")
                if not code:
                    continue
                key = code_key(code)
                if key in products:
                    raise ValueError(f"Trùng mã sản phẩm không phân biệt hoa/thường: {code}")
                products[key] = {
                    field: values[index] if index < len(values) else None
                    for field, index in positions.items()
                }
                products[key]["code"] = code
            if products:
                return products
    finally:
        workbook.close()
    raise ValueError("Không tìm thấy sheet catalogue có cột code và dữ liệu sản phẩm.")


def read_inventory(path: Path) -> list[InventoryRow]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    required = {"ma hang", "ten hang", "mo ta", "ma kho", "cuoi ky"}
    try:
        for worksheet in workbook.worksheets:
            try:
                header_row, positions = locate_headers(worksheet, required)
            except ValueError:
                continue
            result: list[InventoryRow] = []
            for source_row, values in enumerate(worksheet.iter_rows(min_row=header_row + 2, values_only=True), start=header_row + 2):
                def get(header: str) -> object:
                    index = positions[header]
                    return values[index] if index < len(values) else None

                internal_code = text(get("ma hang"))
                if not internal_code or normalized_header(internal_code).startswith("tong cong"):
                    continue
                result.append(InventoryRow(
                    source_row=source_row,
                    internal_code=internal_code,
                    item_name=text(get("ten hang")),
                    description=text(get("mo ta")),
                    stock_location=text(get("ma kho")),
                    quantity=parse_number(get("cuoi ky")),
                ))
            return result
    finally:
        workbook.close()
    raise ValueError("Không tìm thấy sheet tồn kho đúng cấu trúc MISA.")


def build_code_indexes(products: dict[str, dict[str, object]]) -> tuple[dict[str, set[str]], dict[str, set[str]]]:
    exact: dict[str, set[str]] = defaultdict(set)
    normalized: dict[str, set[str]] = defaultdict(set)
    for product in products.values():
        code = text(product["code"])
        for alias in [code, *(text(part) for part in code.split("/") if text(part))]:
            exact[code_key(alias)].add(code)
            normalized[normalized_code(alias)].add(code)
    return exact, normalized


def resolve_value(value: str, exact: dict[str, set[str]], normalized: dict[str, set[str]]) -> tuple[str, str, str]:
    """Return product code, rule and a reason when a value is unsafe."""
    for candidate in candidates(value):
        direct = exact.get(code_key(candidate), set())
        if len(direct) == 1:
            return next(iter(direct)), "Mã hàng" if value == candidate else "Biến thể mã hàng", ""
        if len(direct) > 1:
            return "", "ambiguous", f"Candidate {candidate} khớp nhiều mã: {', '.join(sorted(direct))}"
        simplified = normalized.get(normalized_code(candidate), set())
        if len(simplified) == 1:
            return next(iter(simplified)), "Mã hàng chuẩn hóa", ""
        if len(simplified) > 1:
            return "", "ambiguous", f"Candidate chuẩn hóa {candidate} khớp nhiều mã: {', '.join(sorted(simplified))}"
    return "", "unmatched", "Không có mã sản phẩm catalogue tương ứng"


def match_inventory(rows: list[InventoryRow], products: dict[str, dict[str, object]]) -> list[MatchResult]:
    exact, normalized = build_code_indexes(products)
    output: list[MatchResult] = []
    for row in rows:
        product_code, rule, reason = resolve_value(row.internal_code, exact, normalized)
        if not product_code and rule != "ambiguous":
            for alias in extract_alias_codes(row.description):
                product_code, alias_rule, reason = resolve_value(alias, exact, normalized)
                if product_code or alias_rule == "ambiguous":
                    rule = "Alias mô tả" if product_code else alias_rule
                    break
        name, size, packing = extract_name_size_packing(row.item_name, row.internal_code)
        pieces, area = parse_packaging(packing)
        output.append(MatchResult(
            row=row,
            product_code=product_code,
            rule=rule,
            reason=reason,
            name=name,
            size=size,
            packing=packing,
            packing_pcs=pieces,
            packing_m2=area,
        ))
    return output


def group_candidates(matches: list[MatchResult]) -> list[MappingCandidate]:
    grouped: dict[str, list[MatchResult]] = defaultdict(list)
    for match in matches:
        grouped[code_key(match.row.internal_code)].append(match)
    candidates_out: list[MappingCandidate] = []
    for rows in grouped.values():
        internal_code = rows[0].row.internal_code
        matched_codes = {match.product_code for match in rows if match.product_code}
        unsafe = [match for match in rows if match.rule == "ambiguous"]
        if unsafe:
            candidates_out.append(MappingCandidate(internal_code, status="ambiguous", rows=rows, reason=unsafe[0].reason))
        elif len(matched_codes) == 1:
            candidates_out.append(MappingCandidate(internal_code, next(iter(matched_codes)), "matched", rows))
        elif len(matched_codes) > 1:
            candidates_out.append(MappingCandidate(
                internal_code, status="multi_product_conflict", rows=rows,
                reason=f"Cùng mã nội bộ khớp nhiều sản phẩm: {', '.join(sorted(matched_codes))}",
            ))
        else:
            reason = next((match.reason for match in rows if match.reason), "Không khớp")
            candidates_out.append(MappingCandidate(internal_code, status="unmatched", rows=rows, reason=reason))
    return sorted(candidates_out, key=lambda candidate: candidate.internal_code.casefold())


def read_crm_snapshot(path: Path) -> dict[str, str]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    code_aliases = {"product_code", "ma bao gia", "ma san pham", "code"}
    internal_aliases = {"internal_code", "ma noi bo", "ma hang"}
    try:
        for worksheet in workbook.worksheets:
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                positions = header_positions(row)
                code_column = next((positions[name] for name in code_aliases if name in positions), None)
                internal_column = next((positions[name] for name in internal_aliases if name in positions), None)
                if code_column is None or internal_column is None:
                    continue
                existing: dict[str, str] = {}
                for values in worksheet.iter_rows(min_row=row_number + 1, values_only=True):
                    internal_code = text(values[internal_column] if internal_column < len(values) else "")
                    product_code = text(values[code_column] if code_column < len(values) else "")
                    if not internal_code or not product_code:
                        continue
                    key = code_key(internal_code)
                    previous = existing.get(key)
                    if previous and code_key(previous) != code_key(product_code):
                        raise ValueError(f"Snapshot CRM có mã nội bộ {internal_code} gắn nhiều sản phẩm.")
                    existing[key] = product_code
                return existing
    finally:
        workbook.close()
    raise ValueError("Không tìm thấy cột product_code và internal_code trong export CRM.")


def apply_preflight(candidates_in: list[MappingCandidate], existing: dict[str, str]) -> None:
    for candidate in candidates_in:
        if candidate.status != "matched":
            continue
        prior = existing.get(code_key(candidate.internal_code))
        if not prior:
            continue
        if code_key(prior) == code_key(candidate.product_code):
            candidate.status = "already_present"
            candidate.reason = "Mã nội bộ đã được gắn đúng trong CRM"
        else:
            candidate.status = "existing_crm_collision"
            candidate.reason = f"Đã thuộc mã báo giá {prior} trong CRM"


def proposed_updates(candidates_in: list[MappingCandidate], products: dict[str, dict[str, object]]) -> list[dict[str, object]]:
    """Suggest only fields that agree across every matched internal code of a product."""
    source_rows: dict[str, list[tuple[MappingCandidate, MatchResult]]] = defaultdict(list)
    for candidate in candidates_in:
        if candidate.status not in {"matched", "already_present"}:
            continue
        for match in candidate.rows:
            if match.product_code:
                source_rows[candidate.product_code].append((candidate, match))

    proposals: list[dict[str, object]] = []
    for product_code, rows in source_rows.items():
        product = products[code_key(product_code)]
        suggestion: dict[str, object] = {
            "code": product_code,
            "internal_code": ", ".join(sorted({candidate.internal_code for candidate, _ in rows})),
            "rule": ", ".join(sorted({match.rule for _, match in rows})),
        }
        changed: list[str] = []
        conflicts: list[str] = []
        for field in CATALOGUE_FIELDS:
            values = {getattr(match, field) for _, match in rows if getattr(match, field) not in (None, "")}
            if len(values) != 1:
                if len(values) > 1:
                    conflicts.append(field)
                suggestion[field] = ""
                continue
            value = next(iter(values))
            current = product.get(field)
            suggestion[field] = value
            if str(value) != str(current if current is not None else ""):
                changed.append(field)
        if changed:
            suggestion["fields"] = tuple(changed)
            suggestion["changed_fields"] = ",".join(changed)
            suggestion["conflicting_fields"] = ",".join(conflicts)
            proposals.append(suggestion)
    return sorted(proposals, key=lambda row: str(row["code"]).casefold())


def format_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def write_mapping(path: Path, candidates_in: list[MappingCandidate]) -> int:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mapping"
    sheet.append(MAPPING_HEADERS)
    mappings = [candidate for candidate in candidates_in if candidate.status == "matched"]
    for candidate in mappings:
        sheet.append([candidate.product_code, candidate.internal_code])
    format_sheet(sheet, [24, 24])
    workbook.save(path)
    return len(mappings)


def write_audit(
    path: Path,
    candidates_in: list[MappingCandidate],
    proposals: list[dict[str, object]],
    products: dict[str, dict[str, object]],
    catalogue_paths: list[Path],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tóm tắt"
    counts: dict[str, int] = defaultdict(int)
    for candidate in candidates_in:
        counts[candidate.status] += 1
    summary.append(["Hạng mục", "Số lượng"])
    summary.append(["Tổng mã nội bộ nguồn", len(candidates_in)])
    for status in ("matched", "already_present", "existing_crm_collision", "ambiguous", "multi_product_conflict", "unmatched"):
        summary.append([status, counts[status]])
    summary.append(["Catalogue proposals an toàn", len(proposals)])
    summary.append(["File catalogue update đã tạo", len(catalogue_paths)])
    summary.append(["Lưu ý", "Tự kiểm tra các file import trước khi import. Audit này chỉ để đối chiếu, không import vào CRM."])
    format_sheet(summary, [34, 100])

    mapping = workbook.create_sheet("Mapping audit")
    mapping.append(["Mã nội bộ", "Mã báo giá", "Trạng thái", "Lý do", "Quy tắc", "Kho", "Tồn cuối kỳ", "Dòng nguồn", "Tên hàng", "Mô tả"])
    for candidate in candidates_in:
        for match in candidate.rows:
            mapping.append([
                candidate.internal_code, candidate.product_code, candidate.status, candidate.reason,
                match.rule, match.row.stock_location, match.row.quantity, match.row.source_row,
                match.row.item_name, match.row.description,
            ])
    format_sheet(mapping, [22, 20, 25, 45, 22, 15, 16, 12, 50, 50])

    updates = workbook.create_sheet("Catalogue audit")
    headers = ["Mã báo giá", "Mã nội bộ", "Quy tắc", "Các field đề xuất", "Field nguồn mâu thuẫn"]
    for field in CATALOGUE_FIELDS:
        headers.extend([f"Hiện tại: {field}", f"Đề xuất: {field}"])
    updates.append(headers)
    for proposal in proposals:
        product_code = text(proposal["code"])
        product = products[code_key(product_code)]
        row = [
            product_code, proposal["internal_code"], proposal["rule"], proposal["changed_fields"],
            proposal["conflicting_fields"],
        ]
        for field in CATALOGUE_FIELDS:
            row.extend([product.get(field, ""), proposal[field]])
        updates.append(row)
    format_sheet(updates, [20, 20, 22, 28, 28, 28, 28, 28, 28, 22, 22, 22, 22, 22, 22])
    workbook.save(path)


def write_catalogue_updates(output_dir: Path, proposals: list[dict[str, object]], timestamp: str) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    grouped: dict[tuple[str, ...], list[list[object]]] = defaultdict(list)
    for proposal in proposals:
        fields = proposal["fields"]
        if not isinstance(fields, tuple) or not fields:
            raise ValueError(f"Proposal catalogue không có field hợp lệ: {proposal['code']}")
        output_row = [proposal["code"], *(proposal[field] for field in fields)]
        if any(value in (None, "") for value in output_row[1:]):
            raise ValueError(f"Proposal catalogue có giá trị trống: {proposal['code']}")
        grouped[fields].append(output_row)

    paths: list[Path] = []
    for fields in sorted(grouped):
        suffix = "_".join(fields)
        path = output_dir / f"catalogue_update_{timestamp}_{suffix}.xlsx"
        output = Workbook()
        sheet = output.active
        sheet.title = "Catalogue update"
        sheet.append(["code", *fields])
        for row in sorted(grouped[fields], key=lambda row: code_key(row[0])):
            sheet.append(row)
        format_sheet(sheet, [22, *([28] * len(fields))])
        output.save(path)
        paths.append(path)
    return paths


def run(args: argparse.Namespace) -> int:
    catalogue_path = Path(args.catalogue)
    inventory_path = Path(args.inventory)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if not args.crm_internal_codes:
        raise ValueError("Thiếu --crm-internal-codes. Hãy export 'Tồn Kho & Mã Nội Bộ' từ CRM để preflight.")
    products = read_catalogue(catalogue_path)
    matches = match_inventory(read_inventory(inventory_path), products)
    candidates_in = group_candidates(matches)
    apply_preflight(candidates_in, read_crm_snapshot(Path(args.crm_internal_codes)))
    proposals = proposed_updates(candidates_in, products)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    mapping_path = output_dir / f"internal_code_mapping_{timestamp}.xlsx"
    audit_path = output_dir / f"matching_audit_{timestamp}.xlsx"
    mapping_count = write_mapping(mapping_path, candidates_in)
    catalogue_paths = write_catalogue_updates(output_dir, proposals, timestamp)
    write_audit(audit_path, candidates_in, proposals, products, catalogue_paths)
    print(f"Đã tạo Mapping ({mapping_count} mã mới): {mapping_path}")
    if catalogue_paths:
        for path in catalogue_paths:
            print(f"Đã tạo Catalogue update: {path}")
    else:
        print("Không có field catalogue đủ điều kiện để cập nhật.")
    print(f"Đã tạo file audit: {audit_path}")
    print("Tự kiểm tra output, rồi import Mapping và từng file Catalogue update vào CRM.")
    return 0


def parse_args() -> argparse.Namespace:
    base_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="Tạo file import mã nội bộ và catalogue, kèm audit đối chiếu cho CRM.")
    parser.add_argument("--catalogue", default=str(base_dir / DEFAULT_CATALOGUE), help="File catalogue export từ CRM")
    parser.add_argument("--inventory", default=str(base_dir / DEFAULT_INVENTORY), help="File tổng hợp tồn kho")
    parser.add_argument(
        "--crm-internal-codes",
        default=str(base_dir / "crm_internal_codes.xlsx"),
        help="Export Tồn Kho & Mã Nội Bộ mới nhất từ CRM",
    )
    parser.add_argument("--output-dir", default=str(base_dir / "output"), help="Thư mục ghi workbook")
    return parser.parse_args()


if __name__ == "__main__":
    try:
        raise SystemExit(run(parse_args()))
    except (ValueError, FileNotFoundError) as error:
        print(f"Error: {error}", file=sys.stderr)
        raise SystemExit(2)
