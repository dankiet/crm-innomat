from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import generate_import_files as generator


class GeneratorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.catalogue = self.root / "catalogue.xlsx"
        self.inventory = self.root / "inventory.xlsx"
        self.snapshot = self.root / "crm.xlsx"
        self.output = self.root / "output"
        self._write_catalogue()
        self._write_inventory()
        self._write_snapshot()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _write_catalogue(self) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "San pham"
        sheet.append(["id", "code", "name", "size", "packing", "packing_m2", "packing_pcs"])
        sheet.append([1, "SP-01", "Tên cũ", "100x100", "", "", ""])
        sheet.append([2, "SP-02", "Sản phẩm 2", "", "", "", ""])
        workbook.save(self.catalogue)

    def _write_inventory(self) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "TỔNG HỢP TỒN KHO"
        sheet.append(["Báo cáo tồn kho"])
        sheet.append([])
        sheet.append([])
        sheet.append(["Mã hàng", "Tên hàng", "Mô tả", "Mã kho", "Đầu kỳ", "Nhập", "Xuất", "Cuối kỳ"])
        sheet.append(["", "", "", "", "", "", "", ""])
        sheet.append(["1-SP-01", "Tên mới, 1-SP-01, 200x200 (10 viên/1.2m²)", "", "KHOQ9", 0, 0, 0, 12])
        sheet.append(["1-SP-01", "Tên mới, 1-SP-01, 200x200 (10 viên/1.2m²)", "", "KHOVP", 0, 0, 0, 5])
        sheet.append(["HANG-ALIAS", "Tên alias, HANG-ALIAS, 300x300 (8 viên/0.8m²)", "SP-02 (HANG-ALIAS)", "KHOQ9", 0, 0, 0, 3])
        sheet.append(["Tổng cộng", "", "", "", "", "", "", 20])
        workbook.save(self.inventory)

    def _write_snapshot(self) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["id", "product_code", "product_name", "internal_code", "stock_location", "quantity_stock"])
        sheet.append([1, "SP-01", "Tên cũ", "EXISTING", "KHOQ9", 1])
        workbook.save(self.snapshot)

    def test_candidates_preserve_direct_match_before_variants(self) -> None:
        result = generator.candidates("1-SP-01")
        self.assertEqual(result[0], "1-SP-01")
        self.assertIn("SP-01", result)

    def test_packaging_decimal_is_not_treated_as_thousands(self) -> None:
        self.assertEqual(generator.parse_packaging("10 viên/1.202m²"), (10.0, 1.202))

    def test_matches_alias_and_deduplicates_multiple_locations(self) -> None:
        products = generator.read_catalogue(self.catalogue)
        matches = generator.match_inventory(generator.read_inventory(self.inventory), products)
        candidates = generator.group_candidates(matches)

        self.assertEqual(len(candidates), 2)
        direct = next(candidate for candidate in candidates if candidate.internal_code == "1-SP-01")
        alias = next(candidate for candidate in candidates if candidate.internal_code == "HANG-ALIAS")
        self.assertEqual(direct.product_code, "SP-01")
        self.assertEqual(len(direct.rows), 2)
        self.assertEqual(alias.product_code, "SP-02")
        self.assertEqual(alias.rows[0].rule, "Alias mô tả")

    def test_mapping_workbook_has_exactly_two_import_columns(self) -> None:
        products = generator.read_catalogue(self.catalogue)
        candidates = generator.group_candidates(generator.match_inventory(generator.read_inventory(self.inventory), products))
        generator.apply_preflight(candidates, generator.read_crm_snapshot(self.snapshot))
        path = self.root / "mapping.xlsx"
        count = generator.write_mapping(path, candidates)

        workbook = openpyxl.load_workbook(path, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            self.assertEqual(sheet.title, "Mapping")
            self.assertEqual(tuple(cell.value for cell in sheet[1]), generator.MAPPING_HEADERS)
            self.assertEqual(sheet.max_column, 2)
            self.assertEqual(count, 2)
        finally:
            workbook.close()

    def test_preflight_blocks_internal_code_owned_by_another_product(self) -> None:
        products = generator.read_catalogue(self.catalogue)
        candidates = generator.group_candidates(generator.match_inventory(generator.read_inventory(self.inventory), products))
        existing = generator.read_crm_snapshot(self.snapshot)
        existing[generator.code_key("HANG-ALIAS")] = "SP-01"
        generator.apply_preflight(candidates, existing)

        conflict = next(candidate for candidate in candidates if candidate.internal_code == "HANG-ALIAS")
        self.assertEqual(conflict.status, "existing_crm_collision")

    def test_catalogue_updates_are_created_directly_by_field_subset(self) -> None:
        proposals = [
            {"code": "SP-02", "fields": ("name", "size"), "name": "Tên 2", "size": "200x200"},
            {"code": "SP-01", "fields": ("name", "size"), "name": "Tên 1", "size": "100x100"},
            {"code": "SP-03", "fields": ("packing", "packing_m2"), "packing": "8 viên/0.8m²", "packing_m2": 0.8},
        ]

        paths = generator.write_catalogue_updates(self.output, proposals, "test")
        self.assertEqual(len(paths), 2)
        workbooks = {}
        for path in paths:
            workbook = openpyxl.load_workbook(path, data_only=True)
            sheet = workbook.worksheets[0]
            workbooks[tuple(cell.value for cell in sheet[1])] = [
                tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2)
            ]
            workbook.close()
        self.assertEqual(workbooks[("code", "name", "size")], [("SP-01", "Tên 1", "100x100"), ("SP-02", "Tên 2", "200x200")])
        self.assertEqual(workbooks[("code", "packing", "packing_m2")], [("SP-03", "8 viên/0.8m²", 0.8)])

    def test_catalogue_updates_skip_empty_proposals(self) -> None:
        self.assertEqual(generator.write_catalogue_updates(self.output, [], "test"), [])

    def test_proposals_exclude_conflicting_source_fields(self) -> None:
        products = generator.read_catalogue(self.catalogue)
        matches = generator.match_inventory(generator.read_inventory(self.inventory), products)
        matches[1].name = "Tên khác"
        candidates = generator.group_candidates(matches)
        proposals = generator.proposed_updates(candidates, products)

        proposal = next(item for item in proposals if item["code"] == "SP-01")
        self.assertNotIn("name", proposal["fields"])
        self.assertIn("name", proposal["conflicting_fields"])
        self.assertIn("size", proposal["fields"])


if __name__ == "__main__":
    unittest.main()
